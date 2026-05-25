import os
import io
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def generate_excel_report(all_results, chunk_sizes):
    """
    all_results: list of dicts:
        {
            "question": "...",
            "methodologies": {
                "Semantic": {
                    "openai": { 128: {...}, 256: {...}, ... },
                    "google": { 128: {...}, ... }
                },
                ...
            }
        }
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    highlight_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    header_font = Font(bold=True)
    side_header_font = Font(bold=True)
    
    for q_idx, q_data in enumerate(all_results, 1):
        q_text = q_data["question"]
        sheet_name = f"Q{q_idx}"[:31]
        ws = wb.create_sheet(title=sheet_name)
        
        # Write Question Info
        ws.cell(row=1, column=1, value="Question:").font = header_font
        ws.cell(row=1, column=2, value=q_text)
        
        # Build the dynamic list of strategy-model columns
        methodologies = list(q_data.get("methodologies", {}).keys())
        providers = ["openai", "google"]
        provider_names = {
            "openai": "OpenAI (text-embedding-3-large)",
            "google": "Google (gemini-embedding-2)"
        }
        
        # Headers on row 3
        ws.cell(row=3, column=1, value="Chunk Size").font = header_font
        ws.cell(row=3, column=2, value="Metric").font = header_font
        
        col_mapping = [] # list of tuples: (col_idx, strategy, provider)
        curr_col = 3
        for meth in methodologies:
            for prov in providers:
                cell = ws.cell(row=3, column=curr_col, value=f"{meth}\n{provider_names[prov]}")
                cell.font = header_font
                cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
                col_mapping.append((curr_col, meth, prov))
                curr_col += 1
                
        # Populate rows. For each size, we have 4 rows: Precision, Time, Answer, Chunks.
        curr_row = 4
        for size in chunk_sizes:
            # Merged Chunk Size cell in Column A
            cell_size = ws.cell(row=curr_row, column=1, value=f"{size} tokens")
            cell_size.font = side_header_font
            cell_size.alignment = Alignment(horizontal='center', vertical='center')
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row+3, end_column=1)
            
            # Row 1: Precision
            ws.cell(row=curr_row, column=2, value="Precision (Cosine Sim)").font = side_header_font
            best_col = -1
            best_prec = -1.0
            
            for col_idx, meth, prov in col_mapping:
                res = q_data["methodologies"][meth][prov][size]
                val = res["precision_at_3"]
                ws.cell(row=curr_row, column=col_idx, value=val)
                if val > best_prec:
                    best_prec = val
                    best_col = col_idx
            
            # Highlight best precision cell for this chunk size
            if best_col != -1:
                ws.cell(row=curr_row, column=best_col).fill = highlight_fill
                
            # Row 2: Retrieval Time
            curr_row += 1
            ws.cell(row=curr_row, column=2, value="Retrieval Time (s)").font = side_header_font
            for col_idx, meth, prov in col_mapping:
                res = q_data["methodologies"][meth][prov][size]
                val = res.get("elapsed_time", 0.0)
                ws.cell(row=curr_row, column=col_idx, value=val)
                
            # Row 3: Answer
            curr_row += 1
            ws.cell(row=curr_row, column=2, value="Answer").font = side_header_font
            for col_idx, meth, prov in col_mapping:
                res = q_data["methodologies"][meth][prov][size]
                val = res["answer"]
                cell_val = ws.cell(row=curr_row, column=col_idx, value=val)
                cell_val.alignment = Alignment(wrap_text=True, vertical='top')
                
            # Row 4: Retrieved Chunks
            curr_row += 1
            ws.cell(row=curr_row, column=2, value="Retrieved Chunks").font = side_header_font
            for col_idx, meth, prov in col_mapping:
                res = q_data["methodologies"][meth][prov][size]
                chunks_text = "\n\n".join([
                    f"Chunk {i+1} [Sim: {c.get('cosine_similarity', 0.0):.4f} | RRF: {c.get('rrf_score', 0.0):.4f}]: {c['text']}" 
                    for i, c in enumerate(res['chunks'])
                ])
                cell_val = ws.cell(row=curr_row, column=col_idx, value=chunks_text)
                cell_val.alignment = Alignment(wrap_text=True, vertical='top')
                
            curr_row += 1 # advance to next block
            
        # Sizing
        ws.row_dimensions[3].height = 40
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        for col_idx, _, _ in col_mapping:
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 50

    # Save to buffer
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer
